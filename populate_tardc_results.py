from pathlib import Path
from copy import deepcopy

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
BASE = ROOT / "output" / "TARDC_draft_paper_v5_technical_revision.docx"
OUT = ROOT / "output" / "TARDC_draft_paper_v6_with_results.docx"
RESULTS = ROOT / "tardc_experiments" / "results" / "full"
FIGURES = RESULTS / "figures"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths=None):
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[c])
            if r == 0:
                set_cell_shading(cell, "1F4E78")
            elif r % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
                    if r == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def paragraph_element(doc, text, style=None, italic=False, center=False):
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    run.italic = italic
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para._p


def figure_elements(doc, path, caption, width=6.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set("title", caption.split(".", 1)[0])
    shape._inline.docPr.set("descr", caption)
    cap = doc.add_paragraph()
    cap.style = "Caption"
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    cap_run = cap.add_run(caption)
    cap_run.italic = True
    cap_run.font.color.rgb = RGBColor(0, 0, 0)
    return [p._p, cap._p]


def table_elements(doc, caption, headers, rows, widths):
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    for run in cap.runs:
        run.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table, widths)
    return [cap._p, table._tbl]


def remove_between(start, end):
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def insert_before(anchor, elements):
    for element in elements:
        anchor._p.addprevious(element)


doc = Document(BASE)
paras = doc.paragraphs

abstract = (
    "Distributed secondary control removes the single point of failure associated with a centralized microgrid controller, "
    "but exchanged measurements can be manipulated by compromised agents or communication links. This paper proposes "
    "physics-informed trust-adaptive resilient dynamic consensus (TARDC) for frequency restoration, voltage regulation, "
    "and proportional active-power sharing in inverter-based AC microgrids. Locally computable cyber-temporal and physical "
    "evidence drives asymmetric direct trust, trust-gated W-MSR filtering, normalized dynamic aggregation, adaptive event "
    "triggering, and a defined degraded mode. The analysis establishes convex aggregation, finite exclusion under persistent "
    "evidence, a conditional ultimate disagreement bound, and instantaneous feasibility of projected commands. A reduced-order "
    "dynamic benchmark parameterized from a published five-generator microgrid was evaluated using 30 paired held-out seeds, "
    "seven controllers, scenarios S0-S7, and matched no-attack counterfactuals. Across the claimed S0-S6 scope, TARDC achieved "
    "a 100% safe-trial rate and zero false isolation. Its mean frequency RMSE was statistically indistinguishable from fixed-rate "
    "TARDC (difference -1.5e-6 Hz, Holm-adjusted p=0.839), while communication fell by 65.5%. The mean maximum attack-caused "
    "frequency difference was 0.0114 Hz in S1-S6, compared with 0.2706 Hz for ordinary dynamic consensus. Under S7, which "
    "violates the F-local assumption, the safe-trial rate fell to zero and the mean attack-caused difference reached 1.3785 Hz. "
    "These results support a resilience-communication trade-off within the stated model and assumptions, not EMT-level or "
    "hardware-level validation."
)
doc.paragraphs[2].text = abstract

replacements = {
    "A validation design that separates calibration and evaluation, reproduces the original five-generator baseline, tests coordinated stealth attacks and mixed cyber-physical disturbances, and reports control, detection, communication, and computational metrics over paired random seeds.":
        "A reproducible 30-seed evaluation with disjoint calibration, matched no-attack counterfactuals, seven controller variants, S0-S7 cyber-physical scenarios, and scaling tests up to 50 agents. The results quantify both the within-scope resilience-communication trade-off and the beyond-assumption failure boundary.",
    "The experimental ablations in Section 6 are designed to test whether each added mechanism contributes materially.":
        "The ablations in Sections 6 and 7 test whether each added mechanism contributes materially and expose the cost of adaptive communication.",
}
for para in doc.paragraphs:
    for old, new in replacements.items():
        if old in para.text:
            para.text = para.text.replace(old, new)

for para in doc.paragraphs:
    if para.text.startswith("System A reproduces the five-generator"):
        para.text = (
            "System A uses the five-generator low-voltage ratings, loads, droop coefficients, communication interval, and "
            "event sequence reported in [1]. Because the original PSCAD/EMTDC project and complete network/controller data "
            "were not released, the present implementation is an aggregate reduced-order frequency and reactive-voltage "
            "benchmark rather than an electromagnetic-transient reproduction. The published six-link graph is retained as "
            "a reference, and three explicitly reported preauthorized overlay links are added so every node has at least "
            "2F+1 neighbors for F=1. System B scales the same agent-level mechanism to 10, 25, and 50 heterogeneous nodes."
        )
    elif para.text.startswith("Residual scales, evidence weights"):
        para.text = (
            "Residual scales and threshold ranges were fixed from an analytical development sample generated with the declared "
            "sensor-noise model and seed 91337. Evaluation used held-out seeds 0-29. Controller parameters, attack definitions, "
            "topologies, and safety limits were frozen before the reported runs; no result-dependent retuning was performed."
        )
    elif para.text.startswith("Use at least 30 paired random seeds"):
        para.text = (
            "Thirty paired random seeds were used for every stochastic comparison. The analysis reports mean, standard "
            "deviation, median, 95% bootstrap confidence intervals, and paired effects. Friedman tests were followed by "
            "Holm-corrected Wilcoxon signed-rank comparisons. Each attacked run was also paired with the identical physical "
            "realization with the cyber attack disabled."
        )
    elif para.text.startswith("Each run stores the complete frozen configuration"):
        para.text = (
            "Each run stored the frozen configuration, software versions, seed, topology, attack definition, control-rate time "
            "series, and derived metrics. The 1,950 primary traces and 1,650 matched counterfactual traces are separated from "
            "the calibration artifact. Seed 0 was selected in advance for representative time-series figures; all numerical "
            "comparisons use the complete paired seed set."
        )
    elif para.text == "Table 6. Prespecified outputs to be populated after experimental execution.":
        para.text = "Table 6. Prespecified outputs generated by the experimental pipeline."

heading7 = next(p for p in doc.paragraphs if p.text == "7 Results and Discussion")
heading8 = next(p for p in doc.paragraphs if p.text == "8 Limitations")
remove_between(heading7, heading8)

scratch = doc
elements = []
elements.append(paragraph_element(scratch,
    "The complete experiment matrix contained 1,950 primary runs. Every attacked System A run was paired with a no-attack "
    "counterfactual that retained the same controller, physical event sequence, topology, and seed. Therefore, attack-caused "
    "frequency deviation denotes the pointwise difference between paired trajectories rather than a comparison across unequal "
    "physical scenarios. Results are evidence for the released reduced-order model only."))

elements.append(paragraph_element(scratch, "7.1 Frozen Configuration and Nominal Control", "Heading 2"))
table7_rows = [
    ["Simulation", "dt=0.02 s; Ts=0.2 s; duration=20 s; evaluation seeds 0-29"],
    ["DG ratings", "Pmax=[90, 50, 40, 80, 30] kW; Qmax=[72, 40, 32, 64, 24] kVar"],
    ["Loads", "P=[12, 45, 60, 60, 12] kW; Q=[6, 27, 24, 36, 9] kVar"],
    ["Droop coefficients", "mp=[0.0349, 0.0628, 0.0785, 0.0393, 0.1050] Hz/kW; mq=[0.0802, 0.2093, 0.3157, 0.1443, 0.4090] V/kVar"],
    ["Communication graph", "Published edges plus overlay links (1,5), (2,3), and (2,5), using one-based agent labels"],
    ["Trust", "thetaR=0.80; thetaA=2.40; tauQ=0.35; tauR=0.72; eta_d=0.18; eta_r=0.015; HQ=3; HR=5"],
    ["Freshness and event trigger", "Hmin=1; Hrefresh=5; Hmax=8; sigma0=0.018; sigma min/max=0.006/0.040"],
    ["Reduced-order plant", "Frequency inertia=400; damping=28; active-power time constant=0.10 s; voltage inertia=5; damping=4; reactive time constant=0.25 s"],
    ["Safety definition", "49.5-50.5 Hz and 380-420 V throughout the trial"],
    ["Data separation", "Development seed 91337; held-out evaluation seeds 0-29; configuration hash 1f407af2...42cf800"],
]
elements += table_elements(scratch, "Table 7. Frozen electrical, control, communication, trust, and evaluation parameters.",
                           ["Parameter group", "Frozen value"], table7_rows, [1.75, 4.65])
elements.append(paragraph_element(scratch,
    "Under S0, TARDC maintained a 100% safe-trial rate. Mean frequency RMSE was 0.18390 Hz, the mean frequency nadir was "
    "49.51859 Hz, mean voltage RMSE was 1.00039 V, and mean proportional-sharing error was 0.04898. B1 produced a similar "
    "nadir (49.51859 Hz) but a larger sharing error (0.06007) and required 18,180 packets rather than 599.5. Figure 3 shows "
    "the prespecified seed-0 nominal trace. The visible disturbance at 13 s is the one-second island-detection delay inherited "
    "from the source event sequence."))
elements += figure_elements(scratch, FIGURES / "figure3_nominal.png",
    "Figure 3. System A nominal frequency, average voltage, and active-power-sharing response for B1 and TARDC (P), representative held-out seed 0.")

metrics = pd.read_csv(RESULTS / "run_metrics.csv")
within = metrics[(metrics.n == 5) & metrics.scenario.isin([f"S{i}" for i in range(7)])]
table8_rows = []
for c in ["B0", "B1", "B2", "B3", "B4", "B5", "P"]:
    g = within[within.controller == c]
    table8_rows.append([c, f"{g.frequency_rmse_hz.mean():.5f}", f"{g.voltage_rmse_v.mean():.5f}",
                        f"{g.sharing_error.mean():.5f}", f"{g.safe_trial.mean():.3f}",
                        f"{g.packets.mean():.1f}", f"{g.update_time_ms.mean():.3f}"])
elements += table_elements(scratch, "Table 8. Aggregate System A results across the claimed S0-S6 scope and 30 paired seeds.",
    ["Controller", "f RMSE Hz", "V RMSE V", "Sharing error", "Safe rate", "Packets", "Update ms"],
    table8_rows, [0.65, 0.85, 0.85, 1.0, 0.8, 0.85, 0.9])

elements.append(paragraph_element(scratch, "7.2 Attack Resilience and Degraded Operation", "Heading 2"))
elements.append(paragraph_element(scratch,
    "Across S1-S6, TARDC preserved all 180 trials within the declared frequency and voltage limits and produced no false "
    "isolation of normal links. Its mean maximum attack-caused frequency difference relative to the paired counterfactual was "
    "0.01135 Hz, compared with 0.27056 Hz for ordinary dynamic consensus (B2), 0.02319 Hz for W-MSR alone (B3), 0.02684 Hz "
    "for cyber-only trust (B4), and 0.00843 Hz for fixed-rate TARDC (B5). Thus the event-triggered method sharply improved on "
    "B2-B4 but incurred a small resilience cost relative to its fixed-rate counterpart."))
elements.append(paragraph_element(scratch,
    "Persistent physical inconsistency produced mean detection delays of 0.55 s in S1 and S5 and 2.01 s in the slow-ramp S2 "
    "case. Mean trust separation reached 0.628, 0.580, and 0.685, respectively. Replay S3 was rejected after 0.55 s on average "
    "through sequence/freshness checks; because rejected frames do not update numerical trust, its trust-separation statistic "
    "remained zero. Degraded mode was active for 23.3% of S3 and 41.2% of S4 samples. The threshold-aware S6 attack remained "
    "undetected, as designed, but its mean maximum frequency effect was limited to 0.01451 Hz."))
elements += figure_elements(scratch, FIGURES / "figure4_trust_s2.png",
    "Figure 4. Incoming-link trust, fused evidence score, and frequency response under the S2 slow-ramp attack for TARDC, representative held-out seed 0.")
elements += figure_elements(scratch, FIGURES / "figure5_replay_dos.png",
    "Figure 5. TARDC frequency response and degraded-mode transitions under replay S3 and denial-of-service S4, representative held-out seed 0.")
elements += figure_elements(scratch, FIGURES / "figure6_attack_intensity.png",
    "Figure 6. Mean maximum attack-caused frequency difference relative to matched no-attack counterfactuals; error bars are 95% bootstrap confidence intervals over 30 seeds.")
elements.append(paragraph_element(scratch,
    "S7 deliberately placed two coordinated stealth adversaries in neighborhoods designed for F=1. TARDC then produced a "
    "mean maximum paired frequency difference of 1.37846 Hz, a mean nadir of 48.83679 Hz, and a zero safe-trial rate. This is "
    "reported as an empirical failure boundary, not as contradictory evidence to an F-local claim. It also shows that the "
    "local count monitor cannot identify every coordinated value attack when nominal connectivity remains available."))
elements += figure_elements(scratch, FIGURES / "figure7_adversary_limit.png",
    "Figure 7. Safe-trial rate with one F-local adversary in S5 and two coordinated adversaries beyond the assumed limit in S7.")

elements.append(paragraph_element(scratch, "7.3 Communication and Scalability", "Heading 2"))
elements.append(paragraph_element(scratch,
    "Across S0-S6, adaptive triggering reduced TARDC traffic from 1,759.8 to 606.4 packets per trial relative to B5, a 65.5% "
    "reduction. Compared with the repeated-consensus B1 implementation, the reduction was 96.6%. The defensive response was "
    "not free: S7 increased TARDC traffic to 1,266.3 packets, more than twice its S0 value. In S0 scaling tests, mean packet "
    "counts were 599.5, 1,471.6, 3,684.8, and 7,373.8 for 5, 10, 25, and 50 agents. Mean measured update time increased from "
    "0.431 ms at five agents to approximately 2.1 ms at 50 agents on the execution environment. These timings describe the "
    "Python reference implementation and are not real-time hardware benchmarks."))
elements += figure_elements(scratch, FIGURES / "figure8_scalability.png",
    "Figure 8. TARDC packet count, peak interval traffic, and measured mean update time versus network size under S0.")

elements.append(paragraph_element(scratch, "7.4 Ablation and Statistical Interpretation", "Heading 2"))
table9_rows = [
    ["P-B5", "Frequency RMSE, S0-S6", "-0.000002 Hz", "[-0.000071, 0.000066]", "0.839", "No detectable difference"],
    ["P-B5", "Sharing error, S0-S6", "+0.000053", "[0.000026, 0.000081]", "0.0029", "Detectable but very small cost"],
    ["P-B5", "Packets, S0-S6", "-1153.4", "[-1159.9, -1146.7]", "<1.4e-8", "65.5% traffic reduction"],
    ["P-B2", "Attack-caused max |df|, S1-S6", "-0.25921 Hz", "[-0.26703, -0.25002]", "<1.4e-8", "Large reduction"],
    ["P-B3", "Attack-caused max |df|, S1-S6", "-0.01184 Hz", "[-0.01343, -0.01013]", "<1.4e-8", "Reduction"],
    ["P-B4", "Attack-caused max |df|, S1-S6", "-0.01549 Hz", "[-0.01742, -0.01352]", "<1.4e-8", "Reduction"],
    ["P-B5", "Attack-caused max |df|, S1-S6", "+0.00292 Hz", "[0.00161, 0.00425]", "0.0015", "Small fixed-rate advantage"],
]
elements += table_elements(scratch, "Table 9. Prespecified paired ablation comparisons using per-seed scenario means.",
    ["Contrast", "Outcome", "Mean paired difference", "Bootstrap 95% CI", "Holm p", "Interpretation"],
    table9_rows, [0.65, 1.55, 1.05, 1.25, 0.65, 1.25])
elements.append(paragraph_element(scratch,
    "The Friedman tests rejected equality among the seven controllers for the principal outcomes. Table 9 then applies the "
    "prespecified paired Wilcoxon procedure with Holm correction. TARDC and B5 had indistinguishable frequency RMSE, whereas "
    "TARDC used substantially fewer packets. Adaptive triggering caused a statistically detectable but practically small "
    "increase in sharing error and attack-caused deviation relative to B5. Accordingly, the evidence supports a communication-"
    "resilience trade-off rather than unconditional performance superiority."))

insert_before(heading8, elements)

for para in doc.paragraphs:
    if para.text == "The local adversary and graph robustness assumptions exclude a globally compromised neighborhood.":
        para.text = "The reduced-order aggregate plant does not reproduce inverter switching, line electromagnetic transients, protection dynamics, or controller-hardware effects; EMT and controller-hardware-in-the-loop validation remain necessary."
    elif para.text == "A coordinated attacker may remain below all available residual thresholds or falsify the independent measurements on which the physics tests depend.":
        para.text = "The local adversary and graph assumptions exclude neighborhoods with more than F malicious in-neighbors. S7 confirms that coordinated stealth attacks beyond this limit can violate the declared safety envelope."
    elif para.text == "Optional nodal balance evidence requires independently metered branch flows and is unavailable in some installations.":
        para.text = "A coordinated attacker may remain below every available residual threshold or compromise the independent measurements supporting the physics tests. Optional nodal-balance evidence also requires independently metered branch flows that may be unavailable."
    elif para.text == "Instantaneous reference projection does not prove forward invariance of the nonlinear microgrid state.":
        para.text = "Instantaneous reference projection and the disagreement theorem do not establish nonlinear closed-loop stability or forward invariance of the physical state."
    elif para.text == "Simulation evidence should be followed by controller hardware in the loop validation before deployment claims are made.":
        para.text = "The calibration artifact uses the declared analytical noise model rather than field measurements; threshold transfer to a real microgrid therefore requires new development data and prospective validation."
    elif para.text.startswith("The experimental release should archive"):
        para.text = (
            "The accompanying experiment package contains the Python implementation, dependency list, frozen manifest and "
            "hashes, calibration record, 30 evaluation seeds, 1,950 primary traces, 1,650 counterfactual traces, run-level and "
            "aggregate metrics, statistical tests, and scripts that regenerate Figures 3-8."
        )
    elif para.text.startswith("TARDC extends MAS-based distributed AGC"):
        para.text = (
            "TARDC extends MAS-based distributed AGC and AVC from honest-neighbor averaging to evidence-guided resilient "
            "tracking. The formulation combines locally feasible evidence, asymmetric trust, W-MSR filtering, normalized "
            "dynamic aggregation, bounded secondary commands, adaptive communication, and degraded operation. In the released "
            "30-seed reduced-order benchmark, TARDC kept every S0-S6 trial within the declared safety envelope with zero false "
            "isolation and reduced traffic by 65.5% relative to fixed-rate TARDC without a detectable frequency-RMSE difference. "
            "Its paired attack-caused frequency deviation was much smaller than ordinary consensus and the cyber-only and "
            "W-MSR-only ablations, although fixed-rate TARDC retained a small resilience advantage. When S7 violated the F-local "
            "assumption, every TARDC trial became unsafe, establishing an empirical limit rather than a universal guarantee. "
            "The evidence therefore supports a bounded resilience-communication trade-off for this model; EMT, nonlinear "
            "stability, and controller-hardware validation remain future work."
        )

# Restore global typography on paragraphs created or replaced by python-docx.
for para in doc.paragraphs:
    for run in para.runs:
        if para.style.name not in {"Title", "Heading 1", "Heading 2"}:
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")

doc.core_properties.title = "Physics-Informed Trust-Adaptive Dynamic Consensus for Resilient Secondary Control of AC Microgrids"
doc.core_properties.subject = "TARDC reduced-order simulation results and reproducible evaluation"
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
